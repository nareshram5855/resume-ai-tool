import os
import re
import copy
from docx import Document
from docx.shared import Pt


def _find_experience_section(doc: Document) -> int | None:
    """Find the paragraph index where the Experience section starts."""
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip().lower()
        if re.search(r"experience|work\s+history|employment|projects", text):
            style = (para.style.name or "").lower()
            is_heading = "heading" in style or (
                para.runs and all(r.bold for r in para.runs if r.text.strip())
            )
            if is_heading or (len(text) < 50 and text.isupper()):
                return i
    return None


def _find_skills_section(doc: Document) -> int | None:
    """Find the paragraph index where the Skills section starts."""
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip().lower()
        if re.search(r"skills|technical\s+skills|core\s+competencies|technologies", text):
            return i
    return None


def _find_client_paragraph(doc: Document, client_name: str, start_idx: int) -> int | None:
    """Find the paragraph index for a specific client within the experience section."""
    client_lower = client_name.lower()
    for i in range(start_idx, len(doc.paragraphs)):
        text = doc.paragraphs[i].text.strip().lower()
        if client_lower in text:
            return i
    return None


def _find_last_bullet_for_client(doc: Document, client_idx: int) -> int:
    """Find the last bullet point paragraph index for a client block."""
    last_bullet = client_idx
    for i in range(client_idx + 1, len(doc.paragraphs)):
        text = doc.paragraphs[i].text.strip()
        if not text:
            continue
        # If we hit another heading or client header (has a date), stop
        if re.search(
            r"(20\d{2}|19\d{2})\s*[-–—to]+\s*(20\d{2}|19\d{2}|present|current)",
            text,
            re.IGNORECASE,
        ):
            break
        para = doc.paragraphs[i]
        style = (para.style.name or "").lower()
        if "heading" in style:
            break
        last_bullet = i
    return last_bullet


def _clone_paragraph_style(source_para, new_para):
    """Copy formatting from a source paragraph to a new one."""
    if source_para.runs:
        source_run = source_para.runs[0]
        for run in new_para.runs:
            run.font.name = source_run.font.name
            run.font.size = source_run.font.size
            run.font.bold = source_run.font.bold
            run.font.italic = source_run.font.italic
    if source_para.style:
        new_para.style = source_para.style


def _insert_paragraph_after(doc: Document, ref_para, text: str, style=None):
    """Insert a new paragraph after a reference paragraph."""
    new_para = doc.add_paragraph(text)
    if style:
        new_para.style = style
    # Move the new paragraph to after the reference
    ref_para._element.addnext(new_para._element)
    return new_para


def update_resume(original_path: str, analysis: dict) -> str:
    """Update the DOCX resume with missing points and skills."""
    doc = Document(original_path)
    updated_path = original_path.replace(".docx", "_updated.docx")

    # --- Add missing skills ---
    missing_skills = analysis.get("missing_skills", [])
    if missing_skills:
        skills_idx = _find_skills_section(doc)
        if skills_idx is not None:
            # Find the paragraph with skills content (next paragraph after header)
            for i in range(skills_idx + 1, min(skills_idx + 5, len(doc.paragraphs))):
                text = doc.paragraphs[i].text.strip()
                if text:
                    # Append missing skills
                    new_skills_text = ", ".join(missing_skills)
                    current = doc.paragraphs[i]
                    run = current.add_run(f", {new_skills_text}")
                    if current.runs and len(current.runs) > 1:
                        source_run = current.runs[0]
                        run.font.name = source_run.font.name
                        run.font.size = source_run.font.size
                    break

    # --- Add missing experience points ---
    missing_points = analysis.get("missing_points", [])
    if missing_points:
        exp_idx = _find_experience_section(doc)
        if exp_idx is not None:
            # Group missing points by client
            by_client: dict[str, list[str]] = {}
            for mp in missing_points:
                client = mp["client"]
                if client not in by_client:
                    by_client[client] = []
                by_client[client].append(mp["suggested_point"])

            for client, points in by_client.items():
                client_idx = _find_client_paragraph(doc, client, exp_idx)
                if client_idx is None:
                    continue

                last_bullet_idx = _find_last_bullet_for_client(doc, client_idx)
                ref_para = doc.paragraphs[last_bullet_idx]

                # Insert new bullet points after the last bullet
                for point_text in reversed(points):
                    bullet_text = f"• {point_text}"
                    new_para = _insert_paragraph_after(doc, ref_para, bullet_text)
                    _clone_paragraph_style(ref_para, new_para)

    doc.save(updated_path)
    return updated_path
